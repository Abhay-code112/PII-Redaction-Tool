import io
import re
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

from pii_redactor import Policy, Redactor, images
from pii_redactor.images import Line, id_card_rules

needs_ocr = pytest.mark.skipif(not images.available(), reason="OCR packages not installed")


def L(text, y, x0=30, x1=330, h=26):
    return Line(text, (x0, y, x1, y + h))


def labels(found):
    return sorted((f[0].text, f[3]) for f in found)


def test_id_rules_flag_names_numbers_dates_and_cover_the_line_above():
    lines = [L("GovernmentofIndia", 10), L("MERAJKHAN", 60), L("Father:SudhdanKhan", 95),
             L("DOB:12/12/1988", 130), L("294365933461", 200)]
    found = id_card_rules(lines, [])
    got = {(f[0].text, f[3]) for f in found}
    assert ("MERAJKHAN", "PERSON") in got and ("Father:SudhdanKhan", "PERSON") in got
    assert ("DOB:12/12/1988", "DOB") in got and ("294365933461", "AADHAAR") in got
    assert ("GovernmentofIndia", "PERSON") not in got                       # a header is not a name
    assert any(f[0].text == "x" and f[0].box[3] == 60 for f in found)        # strip above MERAJKHAN (Hindi twin)


def test_id_rules_do_nothing_on_ordinary_images():
    assert id_card_rules([L("QUARTERLYREVENUE", 10), L("Total 1,234", 50)], []) == []


def test_address_block_ends_after_the_pin_row():
    lines = [L("Address:sarayshah", 10), L("Poore Durgi", 40), L("212402", 70), L("Pradesh,212402", 70, 400, 600),
             L("help@uidai.gov.in", 200)]
    covered = {f[0].text for f in id_card_rules(lines + [L("Government of India", 0)], []) if f[3] == "ADDRESS"}
    assert covered == {"Address:sarayshah", "Poore Durgi", "212402", "Pradesh,212402"}


def card(lines):
    img = Image.new("RGB", (720, 60 + 55 * len(lines)), "#dbe9f5")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default(size=34)
    except TypeError:                                    # very old Pillow
        font = ImageFont.load_default()
    for i, t in enumerate(lines):
        draw.text((30, 20 + 55 * i), t, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def compact(img_bytes):
    return re.sub(r"\W", "", "".join(l.text for l in images.read_lines(Image.open(io.BytesIO(img_bytes))))).lower()


@needs_ocr
def test_pan_card_text_is_unreadable_after_redaction():
    data = card(["INCOME TAX DEPARTMENT", "NBWPS1951N", "Name", "VISHAL SINGH", "Date of Birth", "06/05/2000"])
    before = compact(data)
    assert "nbwps1951n" in before and "06052000" in before        # OCR can read it to begin with
    out, info = images.redact_image(data, Redactor(Policy(mode="regex")).detect_in_image_text)
    after = compact(out)
    assert info["regions"] >= 3
    for secret in ("nbwps1951n", "06052000", "vishalsingh"):
        assert secret not in after
    assert "incometaxdepartment" in after                          # the rest of the card is untouched


@needs_ocr
def test_image_without_pii_is_returned_byte_for_byte():
    data = card(["Quarterly revenue chart", "Total sales grew"])
    out, info = images.redact_image(data, Redactor(Policy(mode="regex")).detect_in_image_text)
    assert out == data and not info["regions"]


@needs_ocr
def test_docx_image_is_redacted_in_place_and_thumbnail_dropped(tmp_path):
    png = tmp_path / "card.png"
    png.write_bytes(card(["INCOME TAX DEPARTMENT", "NBWPS1951N", "Date of Birth", "06/05/2000"]))
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("scan attached")
    doc.add_picture(str(png), width=Inches(4))
    doc.save(src)
    report = Redactor(Policy(mode="regex")).redact_file(str(src), str(dst))
    media = [r for r in report.images if r["image"].startswith("/word/media")]
    assert media and media[0]["regions"] >= 2
    names = zipfile.ZipFile(dst).namelist()
    assert not any("thumbnail" in n for n in names)
    assert Document(str(dst)).inline_shapes                              # picture still there


def test_images_can_be_switched_off(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    Document().save(src)
    report = Redactor(Policy(mode="regex", images=False)).redact_file(str(src), str(dst))
    assert report.images == [] and report.warnings == []
